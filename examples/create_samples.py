"""예제 엑셀 파일과 HWPX 템플릿을 생성하는 스크립트."""

from pathlib import Path

from openpyxl import Workbook
from hwpx import HwpxDocument


def create_sample_excel():
    """예제 엑셀 파일을 생성한다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "데이터"

    # 헤더
    ws.append(["이름", "직급", "부서", "날짜"])

    # 데이터
    ws.append(["홍길동", "과장", "총무부", "2026-03-06"])
    ws.append(["김철수", "대리", "인사부", "2026-03-06"])
    ws.append(["이영희", "사원", "개발부", "2026-03-07"])

    output_path = Path(__file__).parent / "sample.xlsx"
    wb.save(output_path)
    print(f"예제 엑셀 생성: {output_path}")
    return output_path


def create_sample_hwpx_template():
    """예제 HWPX 템플릿을 생성한다."""
    doc = HwpxDocument.new()

    doc.add_paragraph("인사발령 통보서")
    doc.add_paragraph("")
    doc.add_paragraph("성명: {{이름}}")
    doc.add_paragraph("직급: {{직급}}")
    doc.add_paragraph("소속: {{부서}}")
    doc.add_paragraph("발령일: {{날짜}}")
    doc.add_paragraph("")
    doc.add_paragraph("위 사항을 통보합니다.")

    output_path = Path(__file__).parent / "template.hwpx"
    doc.save_to_path(output_path)
    print(f"예제 템플릿 생성: {output_path}")
    return output_path


if __name__ == "__main__":
    create_sample_excel()
    create_sample_hwpx_template()
    print("\n예제 파일 생성 완료!")
